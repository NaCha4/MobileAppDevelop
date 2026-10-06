from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'ALLTIME_모바일앱_기획서_초안.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.70)
sec.left_margin = sec.right_margin = Inches(.75)
sec.footer_distance = Inches(.30)

for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    s = doc.styles[name]
    s.font.name = '맑은 고딕'
    s.font.color.rgb = RGBColor(0,0,0)
    s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    s.paragraph_format.widow_control = True
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');s.element.get_or_add_pPr().append(snap)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = Pt(15.5)
normal.paragraph_format.space_after = Pt(6)
for name, size in [('Title',26),('Heading 1',18),('Heading 2',12.5),('Heading 3',11.5)]:
    s=doc.styles[name]; s.font.size=Pt(size); s.font.bold=True
    s.paragraph_format.line_spacing=Pt(size*1.3)
    s.paragraph_format.space_before=Pt(10 if name!='Title' else 0)
    s.paragraph_format.space_after=Pt(7)
    s.paragraph_format.keep_with_next=True
doc.styles['Subtitle'].font.size=Pt(12)
doc.styles['Subtitle'].font.italic=False
doc.styles['Subtitle'].paragraph_format.space_after=Pt(13)
for style in doc.styles:
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)

def p(text='', bold=False, small=False):
    x=doc.add_paragraph()
    r=x.add_run(text); r.bold=bold
    if small: r.font.size=Pt(9.5)
    return x

def h(text, level=2): return doc.add_heading(text,level)
def page(title):
    x=h(title,1);x.paragraph_format.page_break_before=True
def bullet(text):
    x=p('• '+text)
    x.paragraph_format.left_indent=Inches(.12)
    x.paragraph_format.first_line_indent=Inches(-.12)
    return x

def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for i,vals in enumerate([headers]+rows):
        row=t.rows[0] if i==0 else t.add_row()
        nr=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(nr)
        if i==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        for j,(cell,text,w) in enumerate(zip(row.cells,vals,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            cb=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'5');edge.set(qn('w:color'),'D9D9D9');cb.append(edge)
            cp.append(cb)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF3' if i==0 else 'FFFFFF');cp.append(shade)
            mar=OxmlElement('w:tcMar')
            for side,n in [('top',90),('bottom',90),('left',110),('right',110)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(n));e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar)
            q=cell.paragraphs[0];q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=Pt(14.5)
            if i==0 or (j==0 and len(text)<20): q.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if headers[j] in ['응답 인원','비율']: q.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=q.add_run(text);r.font.size=Pt(10.5);r.bold=i==0
    end=doc.add_paragraph();end.paragraph_format.space_after=Pt(3);end.paragraph_format.line_spacing=Pt(3)
    return t

def link_para(label, url):
    x=doc.add_paragraph();x.paragraph_format.space_after=Pt(8)
    rel=x.part.relate_to(url,RT.HYPERLINK,is_external=True)
    a=OxmlElement('w:hyperlink');a.set(qn('r:id'),rel)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'24547B');rp.append(color)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'20');rp.append(size)
    r.append(rp);txt=OxmlElement('w:t');txt.text=label;r.append(txt);a.append(r);x._p.append(a)

def ref(n,title,desc,url,label='원문 열기'):
    q=p(f'[{n}] {title}',bold=True);q.paragraph_format.keep_with_next=True
    q=p(desc,small=True);q.paragraph_format.keep_with_next=True
    link_para(label,url)

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('ALLTIME  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='ALLTIME 기숙사 통합 플랫폼 기획서 초안'
doc.core_properties.subject='한양대학교 ERICA 기숙사 생활 지원 모바일 서비스'
doc.core_properties.author=''
doc.core_properties.keywords='ALLTIME, 기숙사, 모바일앱, 서비스 기획, ERICA'

# 1 서비스 개요
doc.add_heading('ALLTIME 기숙사 통합 플랫폼\n기획서 초안',0)
doc.add_paragraph('한양대학교 ERICA 거주생을 위한 모바일 생활 지원 서비스',style='Subtitle')
p('작성 및 자료 확인 기준  2026년 9월 28일',small=True)
h('서비스 이름')
p('ALLTIME 올타임',bold=True)
p('기숙사 생활에 필요한 시설 정보와 행정 업무를 언제든 한곳에서 확인한다는 의미의 가칭이다. 출시 전 유사 서비스명과 상표를 별도로 확인한다.')
h('서비스 목표')
p('기숙사생이 공용시설 상태와 생활 공지, 민원 처리 현황을 한곳에서 확인해 불필요한 이동과 문의를 줄인다.',bold=True)
h('서비스 개요')
p('ALLTIME은 현재 기숙사에 거주하는 학생과 운영자를 연결하는 모바일 웹앱이다. 학생은 본인이 사는 관의 시설 상태, 확인해야 할 공지, 접수한 민원의 진행 상황을 확인하고, 운영자는 같은 시스템에서 공지를 등록하고 요청을 처리한다.')
p('첫 적용 대상은 한양대학교 ERICA 인재관으로 설정한다. 공지 확인과 민원 처리를 실제 운영 흐름에 연결하고, 세탁실 상태 정보가 헛걸음을 줄이는지 검증한 뒤 다른 관과 공용시설로 확장한다. 운영기관의 협조가 확보되기 전에는 모의 데이터로 기능과 사용성을 검증한다.')
table(['구분','초기 기획 범위'],[
['주 사용자','인재관에 현재 거주하는 재학생 중 세탁실을 자주 이용하는 학생'],
['운영 사용자','기숙사 행정 담당자와 시설 담당자'],
['첫 구현','거주 승인, 시설 상태, 공지 확인, 민원 접수와 처리 이력, 관리자 화면'],
['확장 방향','기숙사 커뮤니티, 시설 예약, 설문, 업체 연계 결제와 장비 연동']
],[1.1,5.9])
p('핵심 가설  시설과 행정 정보를 한 화면으로 연결하면 거주생의 생활 불편과 운영자의 반복 응대가 줄어든다. 만족도 향상과 자치 참여 확대는 실증을 통해 확인할 후속 성과다.',small=True)

# 2 현황
page('1 기획 배경과 현재 이용 환경')
h('시설 이용 전에 판단할 정보가 필요하다')
p('기획의 출발점은 공용시설에 직접 가야 사용 가능 여부를 알 수 있다는 이용 경험이다. 공식 시설 안내에는 세탁실, 체력단련실, 독서실과 스터디룸 등이 소개되어 있다. 다만 인재관의 실시간 상태 확인 수단과 시간대별 혼잡 수준은 현장 조사로 확인해야 한다. 앱이 먼저 해결할 문제는 시설 수 부족 자체보다 이용 전 정보 부족이다. [3]')
h('결제 불편은 관과 장비별로 확인해야 한다')
p('행복관 생활안내의 코인세탁기 6대·건조기 3대는 남학생동과 여학생동 각각의 안내이며 인재관 자료가 아니다. 따라서 인재관의 기기 수나 현금 전용 운영을 전제로 설계하지 않는다. 현행 결제수단, 충전 방식, 고장 시 환불 경로를 조사한 뒤 이용방법 안내와 기존 결제 서비스 연결부터 제공한다. [2]')
h('공지는 이미 존재하므로 접근과 확인을 개선한다')
p('창의인재원 홈페이지에는 창의·인재관 및 행복관 공지가 구분되어 있고, HY-in과 학생자치회 채널 등으로 연결된다. 2026학년도 모집요강은 홈페이지 공지 확인과 메일·문자 수신을 안내한다. ALLTIME은 관별 관련 공지, 마감일, 미확인 항목을 모아 보여주는 방식으로 보완한다. 공지 누락이 얼마나 발생하는지는 별도 조사한다. [1][4]')
h('민원 접점을 연결하고 처리 과정을 보여준다')
p('기존 수리·민원 게시판과 답변상태 표시가 있으며, 인터넷 장애와 시설별 문의는 별도 연락처로 안내된다. 시설 예약·대여도 관리실이나 경비실을 거치는 사례가 있다. 접점이 나뉘어 있다는 사실과 처리 지연이 발생한다는 주장은 구분해야 한다. 담당 배정과 처리 이력이 반복 문의를 줄이는지가 검증 대상이다. [3][5]')
table(['시설 후보','현재 안내 또는 필요한 확인','적용 방향'],[
['세탁실','인재관 장비 수와 결제수단 실사','가동 상태와 고장 안내 우선'],
['스터디룸과 독서실','스터디룸 경비실 예약 안내','이용 안내부터, 예약은 후속'],
['체력단련실과 활동실','출입·정원·운영시간 확인','관측 가능한 경우 혼잡도'],
['공동조리실과 대여물품','위치와 대여 담당이 시설별로 다름','사용법과 접수처 연결']
],[1.35,2.9,2.75])
p('시설 목록은 창의인재원 전체의 확장 후보이며 모두 인재관 내부에 있다는 의미는 아니다. [3]',small=True)

# 3 외부근거
page('2 서비스 필요성을 뒷받침하는 근거')
h('생활 갈등을 다룰 비공개 접수 경로')
p('국립한밭대학교의 2024년도 학생생활관 만족도 조사는 2024년 11월 26일부터 29일까지 진행됐으며, 대상 1,324명 중 811명이 참여했다. 강화가 필요한 중점지도 사항에서는 소란·소음 응답이 가장 많았다. [6]')
table(['중점지도 항목','응답 인원','비율'],[
['소란 및 소음','464명','57.2%'],['도난','142명','17.5%'],['관생실 청소상태','69명','8.5%'],['흡연','68명','8.4%'],['룸메이트 관계','61명','7.5%'],['음주','7명','0.9%']
],[4.2,1.4,1.4])
p('이는 사건 발생률이 아니라 학생들이 지도를 강화하길 바라는 항목의 응답 비율이다. 타 대학 결과를 ERICA의 현황으로 일반화하지 않는다. 생활 갈등도 서비스에서 다룰 필요가 있다는 참고 근거로 사용하고, 공개 게시판의 호실 지목 대신 관리자에게 비공개로 신고하는 흐름을 설계한다.',small=True)
h('모바일 행정 지원의 선행 사례')
p('김성완·손병오·조양현의 2013년 연구는 약 250명 규모의 남학생 기숙사에서 모바일 웹 기반 업무 지원 시스템을 한 학기 운영하고 설문으로 유용성을 평가했다. 기숙사 행정의 모바일 제공 가능성을 뒷받침하는 사례지만, 오래된 단일 기관 연구이므로 ALLTIME의 효과나 현재 이용 의향을 입증하지는 않는다. [7]')
h('전국 규모는 실제 공시 수치로 구분한다')
p('교육부의 2025년 10월 대학정보공시 분석에 따르면, 2025년 4월 1일 기준 일반·교육대 193교의 기숙사 수용가능 인원은 344,806명, 수용률은 22.2%다. 전문대 129교는 56,278명, 15.2%로 별도 집계된다. 두 집단의 수용가능 인원 합계는 401,084명이다. [8]')
p('이 수치는 분석 대상 322교의 공급 규모이며 실제 입주자 수나 확보 가능한 고객 수가 아니다. 전체 대학생 수에 일부 대학의 수용률을 곱하지 않는다. 초기 사업성은 전국 인구 추정보다 한 관에서의 이용률, 운영자 참여, 실제 개선 효과로 판단한다.')

# 4 competitor
page('3 경쟁 서비스와 차별화 방향')
p('비교 기준은 공식 공개 기능이다. 앱 실사용 전에는 데이터 정확도, 게시물 활성도, 이용 만족도를 평가한 것으로 표현하지 않는다. ALLTIME 역시 아래 기능을 이미 제공하는 서비스가 아니라 구현할 기획안이다.')
table(['서비스','확인된 역할과 기능','ALLTIME이 집중할 차이'],[
['긱사\n한국사학진흥재단','기숙사 검색, 비용·시설 정보, 관심 기숙사 비교, 피드와 오류 신고. [9]','입사 전 탐색과 비교에 더해 현재 거주 중인 관의 시설 상태와 행정 요청을 연결한다.'],
['에브리타임','학교 인증 커뮤니티, 시간표, 학식, 중고거래와 일부 학교 열람실 정보 등. [10]','학교 소속과 현재 기숙사 거주 자격을 구분하고 행정실의 실제 처리 과정에 연결한다.'],
['메타클럽','제휴 대학·기숙사 세탁실 검색과 QR 결제 등. [11]','결제를 새로 만드는 것보다 도입된 운영사 서비스와의 연계 가능성을 먼저 검토한다.'],
['창의인재원 홈페이지와 기존 창구','공식 공지, 생활안내, 수리·민원 게시판, 시설별 연락처. [3][4][5]','관별 모바일 홈, 확인할 공지, 시설 상태와 본인 민원 이력을 묶는 이용 흐름을 제안한다.']
],[1.35,2.85,2.8])
h('긱사에 대한 평가 수정')
p('공식 자료에 따르면 긱사는 2025년 12월 24일 정식 출시됐으며, Google Play에는 2026년 9월 9일 업데이트가 표시되어 있다. 따라서 업데이트가 없거나 홈페이지 링크만 제공한다는 단정은 제외한다. 앱 업데이트 날짜와 개별 기숙사 정보의 최신성은 별개이므로 후자는 표본 정보를 대조해야 평가할 수 있다. [9]')
h('차별화의 핵심')
p('ALLTIME의 핵심은 거주 자격, 시설 상태, 공식 공지 확인, 민원 진행 이력을 연결하는 것이다. 커뮤니티와 중고거래는 이미 익숙한 기능이므로 초기 차별점으로 과장하지 않는다. 기존 행정 창구와 중복 입력이 생기지 않도록 담당자가 승인한 하나의 처리 절차를 마련한다.')
h('경쟁 비교를 보완할 실사용 조사')
p('동일한 사용자 과제인 “오늘 사용할 세탁기 확인”, “내 관의 중요 공지 찾기”, “수리 요청 진행상태 확인”을 각 서비스에서 수행한다. 지원 여부, 완료시간, 필요한 화면 이동을 기록하고, 정보의 기준일과 공식 처리 주체를 함께 비교한다. 지원 범위 밖의 과제 실패를 곧바로 서비스 품질 저하로 해석하지 않는다.')

# 5 target/personas
page('4 타겟 사용자와 페르소나')
p('1차 타겟은 인재관에 현재 거주하며 세탁실을 주 1회 이상 이용하는 재학생이다. 운영 타겟은 행정실과 시설 담당자이며 실제 도입 의사결정자는 기숙사 운영기관이다. 2026학년도 모집요강의 신입생 창의관 신청 안내를 고려해 인재관 페르소나는 2학년 이상으로 구성한다. [1]')
p('아래 인물과 생활 패턴은 설계를 위한 가상 가설이다. 실제 인터뷰 결과가 아니며 현장 조사 후 수정한다.',small=True)
h('주 페르소나 김도윤 22세 공학계열 3학년')
p('인재관 거주 남학생. 낮에는 수업과 팀 프로젝트를 진행하고 저녁에 운동과 세탁을 해결한다. 세탁 가능한 시간을 짧게 확보하기 때문에 가기 전에 상태를 확인하고 싶어 한다.')
bullet('목표와 불편  공강이나 취침 전 시간을 계획하고 싶지만, 세탁실의 가동 상태를 확인하러 이동하는 일이 번거롭다.')
bullet('필요한 기능  장비별 상태와 갱신 시각, 가동 종료 알림, 고장 표시, 이용방법 안내.')
bullet('대표 과제  세탁실 방문 전 상태를 확인하고, 정보가 오래됐을 때는 확실한 빈 기기로 오해하지 않는다.')
h('보조 페르소나 이준서 21세 경영계열 2학년')
p('인재관 첫 입사자. 시설 이용 규칙과 문의처를 아직 익히지 못했고, 수업 알림과 기숙사 안내가 함께 오면 필요한 내용을 놓칠 수 있다.')
bullet('목표와 불편  퇴사 일정과 점검 공지를 빠뜨리지 않고, 방의 고장을 올바른 곳에 접수하고 싶다.')
bullet('필요한 기능  내 관의 미확인 공지, 마감일 표시, 사진을 첨부하는 간편 민원, 처리 단계 안내.')
bullet('대표 과제  세탁기 QR로 장비 번호가 채워진 신고 화면에 들어가 증상과 사진을 제출한다.')
h('운영 페르소나 박지현 38세 기숙사 행정 담당자')
p('시설 담당자와 거주생 사이의 문의를 조정한다. 같은 시설의 중복 신고와 처리상태 문의가 생긴다는 가정 아래 운영 화면을 설계한다.')
bullet('목표와 불편  요청을 빠뜨리지 않고 적절한 담당자에게 전달하되 새로운 시스템의 이중 입력 부담을 줄이고 싶다.')
bullet('필요한 기능  관별 접수함, 담당 배정, 중복 요청 연결, 상태 변경 기록, 공지 등록과 확인율 집계.')
bullet('대표 과제  동일 세탁기 고장 신고를 묶어 처리하고 학생에게 단계별 진행상태를 공유한다.')
p('확장 타겟은 다른 관의 거주생과 외국인 학생이다. 시범 운영 성과를 확인한 뒤 관별 규정, 언어, 시설 운영 방식의 차이를 반영한다.',small=True)

# 6 scope
page('5 핵심 기능과 구현 우선순위')
p('첫 버전은 학생 화면과 관리자 화면이 연결되는 최소 기능 제품으로 만든다. 시설 상태는 모의·수동 입력 단계와 실제 장비 연동 단계를 명확히 구분한다.')
table(['우선순위','기능','첫 버전의 완료 기준'],[
['P0 필수','거주 자격 승인','로그인 후 관리자가 승인한 참여자만 해당 관에 접근한다. 자격 만료와 퇴사 후 접근 차단을 제공한다.'],
['P0 필수','시설 상태 조회','장비별 가동 중·가동 종료 추정·확인 필요·고장 상태와 갱신 시각, 데이터 출처를 표시한다.'],
['P0 필수','공지와 확인','관별 공지, 중요도, 마감일, 미확인 목록을 제공한다. 학생의 확인 동작을 저장한다.'],
['P0 필수','민원과 진행상태','시설 QR로 위치를 지정하고 증상·사진을 접수한다. 본인 요청의 담당 배정과 처리 이력을 확인한다.'],
['P0 필수','관리자 화면','공지를 등록하고 요청을 배정·처리한다. 학생과 관리자 간 접근권한을 검증한다.'],
['P1 선택','참여와 생활정보','운영 여력에 따라 설문, 분실물·나눔 게시판, 신고·차단 기능을 추가한다.'],
['P2 확장','실제 장비와 업무 연동','장비 API·센서, 혼잡도, 공식 예약, 기존 업체 결제, 학교 인증 연동은 협의와 검증 후 적용한다.']
],[.95,1.4,4.65])
h('알림은 사용자가 행동할 수 있을 때 제공한다')
p('공지 발행, 민원 단계 변경, 등록한 세탁 세션의 종료를 알림 후보로 둔다. 알림 권한을 거부하거나 지원되지 않는 환경에서도 앱 안의 공지함과 처리 이력은 이용할 수 있어야 한다. 푸시 도착을 공지 열람으로 집계하지 않는다.')
h('커뮤니티의 역할과 운영 범위')
p('초기에는 건물 단위 생활정보, 분실물, 나눔에 집중한다. 호실 공개와 특정인 비난을 허용하지 않고 소음·흡연 등 개인 관련 갈등은 비공개 민원으로 보낸다. 중고거래는 인증만으로 안전이 보장되지 않으므로 별도 거래·신고 정책을 갖춘 뒤 확장한다.')
h('설문은 처리 결과까지 연결한다')
p('식단 선호도나 시설 개선 설문은 운영자가 실제 반영할 수 있는 주제로 한정한다. 응답 결과와 반영 여부를 공개해야 참여가 후속 행동으로 이어진다. 매월 실시하는 방식은 운영자가 수용 가능한지 확인한 뒤 결정한다.')

# 7 UX
page('6 화면 구성과 핵심 사용자 흐름')
p('학생용 기본 메뉴는 홈, 시설, 공지, 내 요청으로 구성하고 계정·거주 자격·알림 설정은 프로필에 둔다. 커뮤니티는 후속 버전에서 추가한다. Figma에서는 아래 화면과 오류·빈 상태를 함께 설계한다.')
table(['화면','우선 보여줄 정보','주요 행동'],[
['홈','내 관, 중요 미확인 공지, 즐겨찾는 시설, 진행 중 요청','시설 보기, 공지 확인, 요청 열기'],
['시설 목록','장비 번호와 상태, 갱신 시각, 운영시간','상세 보기와 새로고침'],
['시설 상세','가동 정보의 출처, 사용법, 요금·결제 안내, 고장 여부','내 세션 알림 설정, 고장 신고'],
['공지 상세','공식 작성자, 대상 관, 게시·수정일, 마감일, 원문','읽고 확인, 원문 열기'],
['민원 접수와 상세','위치·증상·사진, 접수번호, 단계별 시각과 답변','제출, 보완, 재문의'],
['관리자','미배정·처리 중·완료 목록, 담당자와 변경 이력','배정, 상태 변경, 답변, 공지 등록']
],[1.25,3.4,2.35])
h('세탁실 방문을 결정하는 흐름')
p('홈 → 세탁실 → 장비 상태와 마지막 갱신 확인 → 이용 가능한지 판단 → 방문. “가동 종료 추정”은 세탁물 수거 여부를 뜻하지 않는다. 데이터가 오래됐거나 연결이 끊기면 “상태 확인 필요”로 바꾸고 가용 장비 수에서 제외한다.')
h('시설 고장을 접수하는 흐름')
p('시설 QR → 로그인 및 거주 승인 확인 → 관·위치·장비 번호 자동 입력 → 증상과 선택 사진 제출 → 접수번호 발급 → 담당 확인과 배정 → 처리 중 → 완료 안내 → 해결 여부 확인 또는 재문의.')
h('공지를 확인하는 흐름')
p('알림 또는 홈의 미확인 공지 → 내 관에 적용되는 공지 상세 → 내용 확인 → 확인 버튼. 운영자는 대상자 대비 확인율을 보고 추가 안내를 판단한다. 확인 버튼은 이해나 준수를 보장하는 지표가 아니다.')
h('화면 설계에서 함께 검증할 상황')
bullet('로그인 전, 거주 승인 대기, 자격 만료, 공지 없음, 센서 단절, 사진 업로드 실패를 설계한다.')
bullet('색상만으로 상태를 구별하지 않고 텍스트와 아이콘을 함께 사용한다.')
bullet('민원 제출 실패 시 내용을 보존하고 다시 제출해도 중복 접수가 발생하지 않게 한다.')

# 8 feasible
page('7 기술 구현과 현실적인 제약')
h('React와 Firebase로 검증 가능한 범위를 만든다')
p('React 기반 반응형 모바일 웹앱을 만들고 필요 시 홈 화면 설치가 가능한 PWA로 구성한다. Firebase Authentication은 로그인, Firestore는 관·시설·공지·요청·상태 이력, Storage는 민원 사진 저장에 사용한다. 관리자 역할과 알림 발송처럼 신뢰가 필요한 작업은 서버 측에서 처리한다. 개발은 모의 데이터와 테스트 계정으로 시작한다.')
h('세탁기 상태는 가동과 사용 가능을 구분한다')
p('제조사·운영사 API가 있으면 우선 검토한다. 전류·전력 센서로는 가동 패턴을 관찰해 종료를 추정할 수 있으나 정확한 잔여시간이나 세탁물 수거 여부까지 알 수 있다고 가정하지 않는다. 이는 계측 가능한 정보의 한계에 따른 설계 판단이다. 코스·시작시각을 입력해 계산하면 “예상”으로 표시하고 실측 오차를 검증한다.')
p('센서는 시설 담당자가 장비 정격과 설치 안전성을 확인한 뒤 소수 기기로 시험한다. 전원 원격 차단은 범위에서 제외한다. 통신 중단과 센서 오류에는 확인 필요 상태로 전환한다. 제조사 기반 장비 운영 서비스가 존재하므로 자체 하드웨어 개발만을 전제로 하지 않는다. [12]')
h('학생증 연동과 NFC는 별도 과제다')
p('학교 소속 인증이 현재 거주 자격을 증명하지는 않는다. 시범 운영에서는 운영자 승인과 만료되는 개인별 초대 절차를 사용한다. 학생증·HY-in·입사정보 자동 연동은 대학의 제공 범위와 승인이 전제다. 웹 NFC는 지원 기기와 데이터 형식에 제약이 있어 학생증 인증을 보장하지 않는다. QR을 기본으로, URL NFC 태그를 화면 바로가기의 보조 수단으로 둔다. [13]')
h('혼잡도와 결제는 운영 조건을 갖춘 뒤 연결한다')
p('헬스장은 출입 연동이 없으면 정확한 현재 인원을 표시하지 않는다. QR 체크인·체크아웃을 시험할 경우 누락을 관찰해 여유·보통·혼잡의 추정 정보로 제공한다. 세탁 결제는 업체 계약, 장비 제어, 정산·환불 책임이 필요하므로 초기에는 결제방법 안내와 기존 서비스 연결로 한정한다. 세탁기 예약은 선점·미사용 문제를 검증한 뒤 검토한다.')
h('웹 푸시는 보조 전달 수단으로 설계한다')
p('FCM 웹 알림에는 HTTPS, 지원 브라우저와 사용자 권한 등이 필요하다. iOS·iPadOS는 홈 화면에 추가된 웹앱 등 지원 조건을 확인해야 한다. Android와 iPhone 실제 기기에서 알림 허용·거부·백그라운드 상황을 테스트하고 앱 내 미확인함을 유지한다. [14][15]')

# 9 validation
page('8 솔루션 검증 계획과 성공 기준')
p('아래 인원과 목표값은 연구 결과가 아니라 시범 운영을 위한 제안이다. 참여자 모집과 사전 측정 후 운영자와 기준을 확정하고, 달성된 것처럼 표현하지 않는다.')
h('단계별 검증 방법')
bullet('문제 검증  거주생 8~12명과 운영자 2~3명을 인터뷰한다. 세탁 빈도, 헛걸음, 공지 누락, 민원 경로를 최근 2주 경험으로 조사하고 현장 장비·결제수단을 확인한다.')
bullet('사용성 검증  Figma와 기능 시제품을 거주생 5~8명이 사용한다. 시설 상태 확인, 중요 공지 확인, 고장 접수, 처리상태 찾기의 성공 여부와 시간을 측정한다.')
bullet('현장 검증  협조 확보 시 거주생 20~30명을 모집해 도입 전 2주와 후 2주를 비교한다. 가능한 한 같은 참여자와 시간대를 사용하고 시험기간·기기 고장·시설 증설 등 변화 요인을 기록한다.')
table(['검증 항목','측정 정의','예비 통과 기준'],[
['시설 이용 편의','헛걸음 횟수 ÷ 세탁실 이용 시도 횟수. 전후 이용일지와 현장 확인으로 기록','도입 전보다 상대적으로 30% 이상 감소'],
['상태 정보 신뢰','정상 수신 표본의 가동 여부 일치율과 전체 운영시간 중 상태 제공 비율을 함께 측정','가동 여부 일치율과 상태 제공률 각각 90% 이상'],
['공지 접근성','전후 동일한 시점·질문으로 마감 공지의 핵심 내용을 확인하지 못한 비율 측정','전보다 20%p 감소. 기초 누락률 20% 미만이면 5% 이하 목표'],
['반복 문의','이미 접수된 건의 진행상태만 묻는 추가 문의 수 ÷ 전체 접수 건수','전보다 상대적으로 20% 이상 감소'],
['신고와 탐색','접수 화면 진입부터 완료까지 중앙값, 핵심 과제 무도움 성공률','신고 60초 이내, 과제 성공률 80% 이상'],
['지속 사용과 운영','참여자별 주간 이용, 사용 의향 5점 척도, 관리자 건당 처리 소요시간','주간 이용 60% 이상, 의향 4점 이상 70%, 운영 시간 증가 없음']
],[1.1,3.75,2.15])
p('공지 확인 버튼의 기록은 도입 후의 보조지표로 따로 집계한다. 전후 비교에는 동일한 질문·과제를 사용하고 난이도와 기한을 맞춘다. 헛걸음·반복 문의의 기준선이 0이거나 발생 건수가 적으면 감소율 판정을 보류하고 건수와 비율을 보고한다.',small=True)
p('표본이 적으면 성공률과 함께 실제 성공 인원, 누락·중도 이탈, 시간 분포를 제시한다. 단순 전후 비교만으로 인과효과를 확정하지 않는다. 가동 종료와 세탁물 수거는 별도 상태이며 센서 정확도와 헛걸음 감소를 각각 검증한다.',small=True)

# 10 execution
page('9 추진 일정과 운영 조건')
p('학기 프로젝트를 위한 8주 개발·검증 예시다. 팀 규모와 학사 일정, 시설 협의 기간에 맞춰 조정한다. 현장 도입이 지연되면 모의 데이터의 사용성 검증을 완료하고 실제 운영 효과는 후속 과제로 남긴다.')
table(['기간','주요 작업','산출물과 다음 단계 조건'],[
['1~2주','현장 실사와 인터뷰, 운영 절차 합의','시설 목록, 현재 업무 흐름, 문제 우선순위, 참여 가능 운영자'],
['3주','Figma 와이어프레임과 사용성 테스트','학생·관리자 핵심 화면, 오류 상태, 개선 목록'],
['4~5주','React·Firebase P0 기능 구현','승인·시설·공지·민원·관리자 동작, 테스트 계정으로 권한 검증'],
['5~6주','협의된 시범군의 도입 전 측정','2주 기준선, 알림 기기 테스트, 데이터 수집 동의'],
['7~8주','도입 후 측정과 운영자 검토','2주 사용 결과, 문제 수정, 확대 또는 보류 판단']
],[.75,2.45,3.8])
h('운영 책임과 데이터 관리')
p('행정 담당자는 공식 공지와 민원 배정, 시설 담당자는 상태 확인과 처리 결과, 개발팀은 시스템 오류 대응을 맡는 역할안을 협의한다. 앱 접수와 기존 게시판을 모두 처리하게 만들지 않고 하나의 담당 절차로 합친다. 업무시간과 예상 응답시간은 실제 운영자가 확정하며 긴급 상황에는 기존 경비실·긴급 연락 경로를 안내한다.')
p('일반 학생은 본인 민원과 소속 관의 허용 정보에만 접근한다. 호실과 연락처는 업무상 필요한 범위로 제한하고 개인 관련 민원과 사진을 커뮤니티에 공개하지 않는다. 학교 비밀번호와 학생증 사진의 상시 수집은 피한다. 거주 자격 만료, 사진 보관기간, 삭제 절차와 담당자 접근 범위는 시범 운영 전에 정한다.')
p('AI 보조 코딩으로 초안을 만들더라도 역할별 접근, 파일 업로드, 중복 접수, 거주 만료, 장비 연결 실패를 검토한다. Firebase 규칙과 서버 검증을 통과한 테스트 환경에서 확인한 뒤 운영 도입을 협의한다. 실제 운영 데이터 변경, 배포, 결제 연결은 별도 승인 대상이다.')
h('시작 전에 답해야 할 질문')
bullet('인재관 세탁기·건조기는 현재 몇 대이며 실제 결제수단과 상태 확인 방법은 무엇인가?')
bullet('어떤 공지가 어느 채널로 전달되고, 학생은 무엇을 놓치는가?')
bullet('민원은 누가 배정하고 어느 단계에서 학생에게 결과를 알려주는가?')
bullet('기존 운영사와 계약·API·장비 설치 조건을 협의할 수 있는가?')
bullet('운영자와 시범 참여자를 확보하고 도입 전후의 동일 지표를 측정할 수 있는가?')

# 11 corrections
page('10 초안의 주요 주장 검증과 수정')
p('확인된 수치, 현장 조사 가설, 후속 협의 항목을 구분해 발표와 개발 과정에서 같은 기준을 사용한다.')
table(['초기 주장','검증 결과와 반영 방향'],[
['인재관 거주 인원 519명','공식 안내와 2026학년도 자료의 수용인원이다. 실제 입사 인원이 아니다. 시범 모집 규모는 별도 확인한다. [1]'],
['인재관 세탁기 6대와 건조기 3대','공식 확인 수치는 행복관 남·여 학생동 각각의 안내다. 인재관 수치는 현장 조사로 전환한다. [2]'],
['세탁기 이용 시 현금 필수','현재 인재관의 현금 전용 여부를 근거로 삼지 않는다. 코인 안내만으로 다른 결제수단이 없다고 단정하지 않는다. [2]'],
['공지 전용 플랫폼이 필요하다','기존 전용 홈페이지와 메일·문자 안내가 있다. 관련 공지의 모바일 접근과 확인 관리 개선으로 수정한다. [1][4]'],
['민원 방식이 분산되어 확인하기 어렵다','게시판과 답변 상태 표시가 이미 있다. 종류별 접점은 확인되지만 실제 처리 지연은 조사해야 한다. 진행정보 개선 가설로 수정한다. [5]'],
['한밭대 소음·도난 등이 문제다','2024년 중점지도 강화 희망 비율을 사용한다. 사건 발생률이나 ERICA 실태로 표현하지 않는다. [6]'],
['331만명과 22%로 고객 68만명','산술상 331만 × 22%는 72만 8,200명이다. 모집단도 달라 계산을 채택하지 않는다. 2025년 분석 322교 수용가능 401,084명과 실제 고객을 구분한다. [8]'],
['긱사는 업데이트와 기능이 없다','공식 기능과 2026년 업데이트가 확인된다. 데이터 최신성·활성도 평가는 별도 실사용 조사 후 판단한다. [9]'],
['센서로 남은 시간과 빈 기기 확인','전력 정보만으로 정확한 잔여시간·세탁물 수거를 보장하지 않는다. 가동 종료 추정과 확인 필요 상태를 사용한다.'],
['인증된 사용자라 안전한 데이터 확보','거주 인증과 거래 안전·개인정보 보호는 별개다. 동의한 참여자의 이용성 검증과 최소한의 집계로 수정한다.']
],[2.0,5.0])
p('인재관 정원은 공식 페이지 간 상충 표기가 있어 연도가 명시된 2026학년도 자료를 우선한다. 창의인재원 전체 2,743명 역시 수용 규모이며 활성 사용자 수가 아니다. 초기 검증 대상은 한 관의 동의한 참여자로 한정한다. [1]',small=True)

# 12 references
page('11 참고자료와 출처')
p('모든 웹 자료 확인일은 2026년 9월 28일이다. 본문의 대괄호 번호로 연결한다. 게시연도 미표시 페이지의 기기 수·가격·운영시간은 현행 여부를 운영기관에 재확인한다.',small=True)
ref(1,'한양대학교 ERICA 2026학년도 모집요강과 창의인재원 소개','모집요강 인쇄면 49쪽과 51쪽. 인재관 수용인원 519명, 신입생 신청 관과 공지 수신 안내. 전체 소개에도 519명·총 2,743명 표기.','https://goerica.hanyang.ac.kr/upload/BBS0061/20250602121503H3J54S.PDF','2026학년도 모집요강 PDF')
link_para('창의인재원 소개와 수용 규모','https://hydorm.hanyang.ac.kr/service/content/intro_02/index.do')
ref(2,'창의인재원 행복관 생활안내','게시연도 미표시. 남·여 학생동 각각의 코인세탁기 6대와 건조기 3대 안내. 인재관의 수치로 사용하지 않는다.','https://hydorm.hanyang.ac.kr/service/content/happy_02_01/index.do')
ref(3,'창의인재원 복지 및 부대시설','게시연도 미표시. 스터디룸, 독서실, 활동실, 공동조리실, 대여물품과 담당 창구.','https://hydorm.hanyang.ac.kr/service/content/info/index.do')
ref(4,'창의인재원 공식 홈페이지','관별 공지, 입사 업무와 생활정보 연결 구조. 조사 시 2026년 9월 공지 게시 확인.','https://hydorm.hanyang.ac.kr/')
ref(5,'창의인재원 수리 민원접수와 생활안내','공식 게시판 및 2023년 과거 목록의 답변상태. 생활안내에는 별도 장애 연락처가 있다. 과거 목록으로 현재 처리속도를 추정하지 않는다.','https://hydorm.hanyang.ac.kr/service/board/qna/index.do?bdSeq=670&page=8','수리 민원접수 공개 목록')
link_para('창의관과 인재관 생활안내','https://hydorm.hanyang.ac.kr/service/content/info_03/index.do')
ref(6,'국립한밭대학교 2024년도 학생생활관 만족도 설문조사 결과','PDF 2쪽 조사 개요, 4쪽 강화가 필요한 중점지도 사항. 811명 응답. 다른 대학의 참고 사례로 한정한다.','https://dorm.hanbat.ac.kr/attach/filedownloads/do_down/no/1984','조사 결과 원문 PDF')

# 13 refs
page('참고자료 학술 통계와 경쟁 서비스')
ref(7,'김성완 손병오 조양현 2013 기숙사 업무 지원을 위한 모바일 웹 서비스 설계 및 구현','한국정보통신학회논문지 17권 6호, 1389~1398쪽. DOI 10.6109/jkiice.2013.17.6.1389. KCI 서지정보와 초록 기준.','https://www.kci.go.kr/kciportal/ci/sereArticleSearch/ciSereArtiView.kci?sereArticleSearchBean.artiId=ART001780842','KCI 논문 정보와 초록')
ref(8,'교육부 2025년 10월 대학정보공시 분석 결과','2025년 10월 31일 발표. 붙임 1·2 기숙사 수용 현황의 합산값은 401,084명. 일반·교육대 분모는 학부 주·야간 및 대학원 주간 재학생이다. 사이버대·폴리텍·대학원대학 등 86교는 분석에서 제외된다.','https://www.moe.go.kr/boardCnts/viewRenew.do?boardID=294&boardSeq=104279&lev=0&m=020402&opType=N&page=1&s=moe&searchType=null&statusYN=W','교육부 보도자료와 첨부 원문')
ref(9,'한국사학진흥재단 긱사 공식 출시 안내와 앱 정보','2025년 12월 24일 출시 보도자료 및 Google Play. 검색·비교·피드 기능과 앱 업데이트 2026년 9월 9일 확인.','https://play.google.com/store/apps/details?hl=ko&id=kr.co.dormitoryAll','긱사 공식 Google Play')
link_para('한국사학진흥재단 긱사 출시 보도자료','https://www.kasfo.or.kr/kasfo/920/subview.do?enc=Zm5jdDF8QEB8JTJGYmJzJTJGa2FzZm8lMkYzMTUlMkY3NTU1MiUyRmFydGNsVmlldy5kbyUzRnBhc3N3b3JkJTNEJTI2cmdzQmduZGVTdHIlM0QlMjZiYnNDbFNlcSUzRCUyNnJnc0VuZGRlU3RyJTNEJTI2aXNWaWV3TWluZSUzRGZhbHNlJTI2YmJzT3BlbldyZFNlcSUzRCUyNnNyY2hDb2x1bW4lM0QlMjZzcmNoV3JkJTNEJTI2')
ref(10,'에브리타임 공식 앱 소개','Google Play 개발자 제공 설명. 학교 인증 커뮤니티·중고거래·학식·일부 학교 열람실 정보 등 확인.','https://play.google.com/store/apps/details?hl=ko&id=com.everytime.v2')
ref(11,'메타클럽 공식 앱 소개','Google Play 개발자 제공 설명. 제휴 대학·기숙사의 세탁실 검색과 QR 결제 기능 확인.','https://play.google.com/store/apps/details?hl=ko&id=com.fingerverse.metapoint')
ref(12,'LG 상업용 세탁가전 사업 관련 공식 보도자료','2025년 8월 25일. 런드리크루의 장비 원격 운영·진단 등 사례. 인재관 장비와의 연동 가능성이 확인된 것은 아니다.','https://www.lg.co.kr/media/release/29287')

# 14 tech references with research instrument avoids sparse page
page('참고자료 기술 조건과 현장 조사 문항')
ref(13,'Chrome for Developers Web NFC','Web NFC의 지원 환경과 NDEF 범위, 저수준 카드 통신 제한. 학생증 인증 자동 연동의 근거로 사용할 수 없다.','https://developer.chrome.com/docs/capabilities/nfc')
ref(14,'Firebase 웹 Cloud Messaging 시작하기','FCM 웹 알림의 HTTPS, 지원 브라우저, 서비스 워커 및 알림 권한 조건. 구현 시 최신 지원 환경을 재확인한다.','https://firebase.google.com/docs/cloud-messaging/web/get-started')
ref(15,'WebKit Web Push for Web Apps on iOS and iPadOS','2023년 2월 16일. iOS·iPadOS 16.4의 홈 화면 웹앱 알림 지원 소개. 실제 대상 기기에서 권한과 동작 조건을 검증한다.','https://webkit.org/blog/13878/web-push-for-web-apps-on-ios-and-ipados/')
h('거주생 인터뷰 문항')
bullet('최근 2주 동안 세탁실을 몇 번 이용하려 했으며, 기기가 모두 사용 중이어서 돌아온 적은 몇 번인가?')
bullet('최근 확인이 늦었던 기숙사 공지는 무엇이며, 어디에서 처음 확인했는가?')
bullet('최근 시설 문제를 어느 경로로 접수했으며, 접수 후 무엇을 알 수 있었는가?')
bullet('시설 상태·공지·민원 중 가장 필요한 하나는 무엇이며, 현재 대안보다 나아지려면 무엇이 달라져야 하는가?')
bullet('앱이 가동 종료를 표시해도 세탁물이 남아 있을 때 어떤 안내가 있어야 판단에 도움이 되는가?')
h('운영자 인터뷰 문항')
bullet('접수부터 배정·처리·완료 통보까지 단계별 담당자는 누구이며, 반복 문의는 어느 단계에서 발생하는가?')
bullet('공지와 민원 데이터를 제공하거나 앱에서 처리할 수 있는 범위는 어디까지인가?')
bullet('기존 업무를 추가하지 않고 통합하려면 어떤 시스템 또는 업체 협조가 필요한가?')
bullet('시설 상태 확인과 서비스 운영에 매일 투입할 수 있는 시간은 어느 정도인가?')
p('조사 기록에는 응답자의 실명 대신 연구용 식별자를 사용하고, 수집 목적과 보관기간을 설명한 뒤 동의를 받는다. 표본과 원자료를 확보하기 전에는 인터뷰 결과나 효과 수치를 작성하지 않는다.',small=True)

doc.save(OUT)
print(str(OUT))
print(f'Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}')
